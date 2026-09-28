"""采购工具 · 导入价格。

只做一件事：把采购带价格的 BOM（成本核算格式：汇总页 + 每个成品一个 Sheet）里的
**原材料单价**提取出来，按确认的价格日期绑定到物料库的物料上（cost_material_price）。

规则（用户 2026-09-27 定）：
- 只取原材料（在文件里从不作为上级出现的物料）的单价；
- 特例：某个半成品下面所有物料的单价都是 0（或为空），导入这个半成品自己的单价；
- 价格日期由导入人确认（默认从订单号里取）；每个物料可有多条不同日期的价格，
  **同一日期且同一价格**则跳过，同日期不同价格照常新增；
- 不再写半成品/成品的推算价（bom_calc）——部件价格改为按研发 BOM 实时计算。

价格按「去掉版本号的编码」绑定（14ST02001-A01 与 -A02 共用一套价格），与物料卡片价格区同一规则。
表格解析复用研发成本导入（services.rd.cost_import）的列识别与汇总页解析。
"""
import io
from datetime import date, datetime

from openpyxl import load_workbook

from database.base import db
from database.models.product.import_raw import ImportProductRaw
from database.models.rd.cost import CostBomNode, CostMaterialPrice, CostSnapshot, CostSnapshotSku
from result import Result
from services.rd.cost_import import (
    _is_summary_sheet, _parse_sku_sheet, _parse_summary_sheet, _strip_version,
    _suggest_date, _to_float, load_col_aliases,
)
from upload_validation import UploadValidationError

IMPORT_NOTE = '采购导入价格'


def _price(value):
    """单价转成 4 位小数；空/非数字/<=0 视为无价格返回 None。"""
    v = _to_float(value)
    if v is None or v <= 0:
        return None
    return round(v, 4)


def parse_workbook(data: bytes):
    """解析采购 BOM：返回 {order_no, suggested_date, lines:[...], sheets, warnings}。"""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise UploadValidationError('文件无法读取，请确认是 .xlsx 格式') from exc
    try:
        aliases = load_col_aliases()
        order_no = None
        for name in wb.sheetnames:
            if _is_summary_sheet(name):
                order_no, _ = _parse_summary_sheet(wb[name])
                break
        lines, warnings, sheets, finished = [], [], 0, []
        for name in wb.sheetnames:
            if _is_summary_sheet(name):
                continue
            parsed = _parse_sku_sheet(wb[name], aliases)
            if not parsed or not parsed['finished_code']:
                warnings.append(f'Sheet「{name}」无法识别成品品号，已跳过')
                continue
            sheets += 1
            # 订单里包含哪些成品（每个 Sheet 的主件品号）：只记这一层，用于成本历史标注「本产品订单」
            finished.append({'code': parsed['finished_code'], 'name': parsed['finished_name'] or '',
                             'spec': parsed['finished_spec'] or ''})
            for line in parsed['lines']:
                lines.append({**line, 'finished_code': parsed['finished_code']})
        if not lines:
            raise UploadValidationError('文件里没有识别到 BOM 明细行，请确认是采购带价格的 BOM')
        return {
            'order_no': order_no or '', 'suggested_date': _suggest_date(order_no),
            'lines': lines, 'sheets': sheets, 'warnings': warnings, 'finished': finished,
        }
    finally:
        wb.close()


def extract_prices(lines):
    """从 BOM 行里挑出要导入的价格。

    返回 (items, zero_items, conflicts)：
    - items：[{code(基础码), code_with_version, name, spec, price, kind: material|semi}]
    - zero_items：原材料单价为 0/空、无法导入的（特例半成品下面的物料不算，它们由半成品价格覆盖）
    - conflicts：同一物料在文件里出现多个不同单价 → 禁止导入（用户 2026-09-28 定），需先改 Excel
    """
    parents = {_strip_version(l['parent_code']) for l in lines if l.get('parent_code')}
    children_by_parent = {}
    for l in lines:
        children_by_parent.setdefault(_strip_version(l['parent_code']), []).append(l)

    # 部件自身作为子件出现时的正价（外购整件价）
    own_priced = {_strip_version(l['child_code']) for l in lines if _price(l.get('unit_price')) is not None}
    # 特例半成品：父件自身有一条正价采购行，且直接下级单价全为 0/空（Codex 审计 #5：
    # 之前没要求父件自身有价，会把「根本没有价格可导」的部件下级也从缺价提醒里藏掉）
    special = set()
    for parent, kids in children_by_parent.items():
        if parent in own_priced and kids and all(_price(k.get('unit_price')) is None for k in kids):
            special.add(parent)
    covered = set()   # 特例半成品下面的物料（递归），不算「无价格」
    stack = list(special)
    while stack:
        p = stack.pop()
        for k in children_by_parent.get(p, []):
            base = _strip_version(k['child_code'])
            if base not in covered:
                covered.add(base)
                stack.append(base)

    items, zero_items, conflicts = {}, {}, {}
    for l in lines:
        base = _strip_version(l['child_code'])
        if not base:
            continue
        if base in parents and base not in special:
            continue            # 普通部件：价格由下级计算，不导入
        kind = 'semi' if base in special else 'material'
        price = _price(l.get('unit_price'))
        if price is None:
            if kind == 'material' and base not in covered and base not in items:
                zero_items.setdefault(base, {
                    'code': base, 'code_with_version': l['child_code'],
                    'name': l.get('child_name') or '', 'spec': l.get('child_spec') or '',
                })
            continue
        if base in items:
            if items[base]['price'] != price:
                conflicts.setdefault(base, {items[base]['price']}).add(price)
            continue
        zero_items.pop(base, None)
        items[base] = {
            'code': base, 'code_with_version': l['child_code'],
            'name': l.get('child_name') or '', 'spec': l.get('child_spec') or '',
            'category': l.get('child_category') or '', 'price': price, 'kind': kind,
        }
    return (
        list(items.values()), list(zero_items.values()),
        [{'code': c, 'prices': sorted(p)} for c, p in conflicts.items()],
    )


def _parse_date(value):
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value or '').strip(), '%Y-%m-%d').date()
    except ValueError:
        return None


def _erp_base_codes(bases):
    """哪些基础码在 ERP 物料表里有对应物料（按去版本号比较）。"""
    found = set()
    for (code,) in db.session.query(ImportProductRaw.code).all():
        base = _strip_version(code)
        if base in bases:
            found.add(base)
    return found


def _existing_same_day(bases, price_date):
    """这些物料在该日期已有的价格：{(基础码, 价格)}（任何来源，同日同价即视为已存在）。"""
    if not bases or not price_date:
        return set()
    rows = db.session.query(CostBomNode.code, CostMaterialPrice.unit_price).join(
        CostMaterialPrice, CostMaterialPrice.node_id == CostBomNode.id,
    ).filter(
        CostBomNode.code.in_(list(bases)), CostMaterialPrice.price_date == price_date,
    ).all()
    return {(code, round(float(price), 4)) for code, price in rows}


class PurchasePriceImportService:

    def preview(self, data: bytes, price_date=None):
        """解析 + 按日期标出每条价格是新增还是「同日同价」跳过，不写库。"""
        try:
            parsed = parse_workbook(data)
        except UploadValidationError as exc:
            return Result.fail(str(exc))
        items, zero_items, conflicts = extract_prices(parsed['lines'])
        if price_date not in (None, '') and _parse_date(price_date) is None:
            return Result.fail('价格日期格式无效（应为 YYYY-MM-DD）')
        day = _parse_date(price_date) or _parse_date(parsed['suggested_date'])
        existing = _existing_same_day({i['code'] for i in items}, day)
        erp = _erp_base_codes({i['code'] for i in items} | {z['code'] for z in zero_items})
        for i in items:
            i['status'] = 'skip' if (i['code'], i['price']) in existing else 'new'
            i['in_erp'] = i['code'] in erp
        for z in zero_items:
            z['in_erp'] = z['code'] in erp
        return Result.ok(data={
            'order_no': parsed['order_no'],
            'suggested_date': parsed['suggested_date'],
            'price_date': day.isoformat() if day else None,
            'sheets': parsed['sheets'], 'line_count': len(parsed['lines']),
            'items': items,
            'new_count': sum(1 for i in items if i['status'] == 'new'),
            'skip_count': sum(1 for i in items if i['status'] == 'skip'),
            'special_semis': [i for i in items if i['kind'] == 'semi'],
            'zero_items': zero_items, 'conflicts': conflicts,
            'can_import': not conflicts and any(i['status'] == 'new' for i in items),
            'warnings': parsed['warnings'],
        })

    def import_prices(self, data: bytes, price_date, username):
        day = _parse_date(price_date)
        if day is None:
            return Result.fail('请确认价格日期（格式 YYYY-MM-DD）')
        try:
            parsed = parse_workbook(data)
        except UploadValidationError as exc:
            return Result.fail(str(exc))
        items, zero_items, conflicts = extract_prices(parsed['lines'])
        if conflicts:
            detail = '；'.join(f"{c['code']}（{'、'.join(str(p) for p in c['prices'])}）" for c in conflicts[:10])
            return Result.fail(f'文件里有 {len(conflicts)} 个物料出现了多个不同单价，请先在 Excel 里改成一致再导入：{detail}')
        if not items:
            return Result.fail('文件里没有可导入的价格（原材料单价都为 0 或为空）')

        bases = {i['code'] for i in items}
        existing = _existing_same_day(bases, day)
        to_create = [i for i in items if (i['code'], i['price']) not in existing]
        if not to_create:
            # 全部同日同价：不建空批次，避免污染导入记录（Codex 审计 #4）
            return Result.ok(data={
                'batch_id': None, 'order_no': parsed['order_no'], 'price_date': day.isoformat(),
                'created': 0, 'skipped': len(items), 'special_semis': 0,
                'zero_count': len(zero_items), 'conflicts': [],
            }, message='全部价格与该日期已有价格相同，没有新增')
        try:
            nodes = {n.code: n for n in CostBomNode.query.filter(CostBomNode.code.in_(list(bases))).all()}
            batch = CostSnapshot(
                order_no=parsed['order_no'] or None, snapshot_date=day,
                notes=IMPORT_NOTE, created_by=username,
            )
            db.session.add(batch)
            db.session.flush()
            # 只记订单包含的成品（不记 BOM 明细）：成本历史据此区分「本产品订单」和「共用物料变价」
            for f in parsed['finished']:
                db.session.add(CostSnapshotSku(
                    snapshot_id=batch.id, finished_code=f['code'][:64],
                    finished_name=(f['name'] or None) and f['name'][:128],
                    finished_spec=(f['spec'] or None) and f['spec'][:256],
                ))
            created, skipped = 0, len(items) - len(to_create)
            for i in to_create:
                node = nodes.get(i['code'])
                if node is None:
                    node = CostBomNode(
                        code=i['code'], code_with_version=i['code_with_version'] or None,
                        name=(i['name'] or None) and i['name'][:128],
                        spec=(i['spec'] or None) and i['spec'][:512],
                        category=(i['category'] or None) and i['category'][:256],
                        node_type='semi' if i['kind'] == 'semi' else 'material',
                        is_purchased_semi=i['kind'] == 'semi',
                    )
                    db.session.add(node)
                    db.session.flush()
                    nodes[i['code']] = node
                elif i['kind'] == 'semi' and not node.is_purchased_semi:
                    node.is_purchased_semi = True
                db.session.add(CostMaterialPrice(
                    node_id=node.id, unit_price=i['price'], price_date=day,
                    source='bom_import', snapshot_id=batch.id, created_by=username,
                    notes=f"采购导入{('：' + parsed['order_no']) if parsed['order_no'] else ''}",
                ))
                created += 1
            db.session.commit()
        except Exception:
            db.session.rollback()
            raise
        return Result.ok(data={
            'batch_id': batch.id, 'order_no': parsed['order_no'], 'price_date': day.isoformat(),
            'created': created, 'skipped': skipped,
            'special_semis': len([i for i in to_create if i['kind'] == 'semi']),
            'zero_count': len(zero_items), 'conflicts': [],
        })

    @staticmethod
    def history(limit=20):
        """最近的导入记录（只列本功能写入的批次）。"""
        batches = CostSnapshot.query.filter(CostSnapshot.notes == IMPORT_NOTE) \
            .order_by(CostSnapshot.id.desc()).limit(limit).all()
        counts = dict(
            db.session.query(CostMaterialPrice.snapshot_id, db.func.count(CostMaterialPrice.id))
            .filter(CostMaterialPrice.snapshot_id.in_([b.id for b in batches]))
            .group_by(CostMaterialPrice.snapshot_id).all()
        ) if batches else {}
        return Result.ok(data=[{
            'id': b.id, 'order_no': b.order_no,
            'price_date': b.snapshot_date.isoformat() if b.snapshot_date else None,
            'created_by': b.created_by,
            'created_at': b.created_at.strftime('%Y-%m-%d %H:%M') if b.created_at else None,
            'price_count': counts.get(b.id, 0),
        } for b in batches])


purchase_price_import_service = PurchasePriceImportService()

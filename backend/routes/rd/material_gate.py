"""研发物料门禁 HTTP 接口。"""

from io import BytesIO

from flask import g, request
from openpyxl import load_workbook

from auth import is_rd_admin, require_auth
from result import Result
from services.rd.material_gate import material_gate_service
from upload_validation import read_spreadsheet_upload, UploadValidationError

from . import rd_bp


def _permission_denied_response():
    return Result.fail('权限不足：需要研发部管理员权限').to_response(403)


@rd_bp.get('/material-gates')
def list_material_gates():
    return Result.ok(material_gate_service.list_active()).to_response()


@rd_bp.get('/material-gates/all')
@require_auth
def list_all_material_gates():
    if not is_rd_admin():
        return _permission_denied_response()
    return Result.ok(material_gate_service.list_all()).to_response()


@rd_bp.post('/material-gates')
@require_auth
def create_material_gate():
    if not is_rd_admin():
        return _permission_denied_response()
    try:
        username = (getattr(g, 'current_user', {}) or {}).get('username')
        gate = material_gate_service.create(request.get_json(silent=True) or {}, username)
    except ValueError as exc:
        return Result.fail(str(exc)).to_response()
    return Result.ok(gate).to_response()


@rd_bp.put('/material-gates/<int:gate_id>')
@require_auth
def update_material_gate(gate_id):
    if not is_rd_admin():
        return _permission_denied_response()
    try:
        gate = material_gate_service.update(gate_id, request.get_json(silent=True) or {})
    except ValueError as exc:
        return Result.fail(str(exc)).to_response()
    if gate is None:
        return Result.fail('门禁不存在').to_response(404)
    return Result.ok(gate).to_response()


def _set_material_gate_active(gate_id: int, is_active: bool):
    if not is_rd_admin():
        return _permission_denied_response()
    try:
        gate = material_gate_service.set_active(gate_id, is_active)
    except ValueError as exc:
        return Result.fail(str(exc)).to_response()
    if gate is None:
        return Result.fail('门禁不存在').to_response(404)
    return Result.ok(gate).to_response()


@rd_bp.delete('/material-gates/<int:gate_id>')
@require_auth
def delete_material_gate(gate_id):
    if not is_rd_admin():
        return _permission_denied_response()
    if not material_gate_service.delete(gate_id):
        return Result.fail('门禁不存在').to_response(404)
    return Result.ok(None).to_response()


@rd_bp.put('/material-gates/<int:gate_id>/deactivate')
@require_auth
def deactivate_material_gate(gate_id):
    return _set_material_gate_active(gate_id, False)


@rd_bp.put('/material-gates/<int:gate_id>/activate')
@require_auth
def activate_material_gate(gate_id):
    return _set_material_gate_active(gate_id, True)


@rd_bp.post('/material-gates/check')
def check_material_gates():
    payload = request.get_json(silent=True) or {}
    codes = payload.get('codes', [])
    if not isinstance(codes, list):
        return Result.fail('codes 必须是数组').to_response()
    return Result.ok(material_gate_service.check(codes)).to_response()


@rd_bp.post('/material-gates/check-file')
def check_material_gate_file():
    file = request.files.get('file')
    if not file:
        return Result.fail('请上传材料清单文件').to_response()
    if not (file.filename or '').lower().endswith('.xlsx'):
        return Result.fail('仅支持 .xlsx 格式').to_response()
    try:
        content = read_spreadsheet_upload(file, label='材料清单')
    except UploadValidationError as exc:
        return Result.fail(str(exc)).to_response(413 if '不能超过' in str(exc) else 400)

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        first_row = next(rows, None)
        if first_row is None:
            return Result.fail('文件为空').to_response()
        headers = [str(value).strip().lower() if value is not None else '' for value in first_row]
        code_index = next(
            (headers.index(name) for name in ('物料编码', '品号') if name in headers),
            None,
        )
        if code_index is None:
            return Result.fail('未找到"物料编码"或"品号"列').to_response()
        codes = {
            str(row[code_index]).strip()
            for row in rows
            if code_index < len(row) and row[code_index] is not None
            and str(row[code_index]).strip()
        }
    except Exception:
        return Result.fail('文件读取失败').to_response()
    finally:
        if 'workbook' in locals():
            workbook.close()

    result = material_gate_service.check(codes)
    return Result.ok({'scanned_count': len(codes), **result}).to_response()

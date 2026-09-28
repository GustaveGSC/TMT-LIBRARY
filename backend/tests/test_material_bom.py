"""物料 BOM（研发 BOM）：解析校验、单层拆分、ERP 编码对应、展开、反查、补关联与删除保护。"""
import io

import openpyxl
import pytest
from flask import Flask

from database.base import db
import database.models.product.category  # noqa: F401
from database.models.product.finished import ProductFinished
import database.models.product.resource  # noqa: F401
import database.models.product.material_supplier  # noqa: F401
from database.models.rd.cost import CostBomNode, CostMaterialPrice, CostSnapshot, CostSnapshotSku
from database.models.product.erp_code_rules import ErpCodeRule
from database.models.product.import_raw import ImportProductRaw
from database.models.product.material import (
    ErpGroupCategory, MaterialBom, MaterialBomLine, MaterialDisableKeyword, MaterialImage,
    ProductMaterial,
)
from routes.product.material import material_bp
from services.product.material import material_service
from services.product.material_bom import (
    build_single_level_boms, material_bom_service, parse_bom_rows,
)
from upload_validation import UploadValidationError
from utils import now_cst


PDM_HEADERS = ['层次', '名称', '一级分类', '二级分类', '三级分类', '物料编码', '版本',
               '数量', '单位', '描述', '规格', '状态']


def _pdm(rows):
    """rows: (层次, 一级分类, 编码, 版本, 数量, 描述[, 单位])；层次/数量原样写入单元格。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(PDM_HEADERS)
    for row in rows:
        level, cat, code, ver, qty, desc = row[:6]
        unit = row[6] if len(row) > 6 else 'PCS'
        ws.append([level, f'{code}.SLDASM', cat, '01_桌类', None, code, ver, qty, unit,
                   desc, None, '已发布'])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _erp(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['层次', '图号', '品名', '规格', '数量', '单位', '状态'])
    for row in rows:
        ws.append(list(row) + ['已发布'])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _split(data):
    rows, _skipped = parse_bom_rows(data)
    return build_single_level_boms(rows)


# 成品 F（研发 A01）→ 产成品 P + 螺钉 S×4 + 半成品 M；产成品 P → 半成品 M ×2；半成品 M → 原材料 R ×3
# 半成品 M 在成品下又直接挂了一次（内容一致，去重）；14ST10 标准件与带「.」子零件（含其下级）跳过
SAMPLE = [
    ('1',       '11_成品',   'F1', 'A01', 1, '书桌'),
    ('1.1',     '12_产成品', 'P1', 'A01', 1, '桌面'),
    ('1.1.1',   '13_半成品', 'M1', 'A01', 2, '钢架'),
    ('1.1.1.1', '14_原材料', 'R1', 'A01', 3, '方管'),
    ('1.2',     '14_原材料', 'S1', 'A01', 4, '螺钉'),
    ('1.3',     '13_半成品', 'M1', 'A01', 1, '钢架'),
    ('1.3.1',   '14_原材料', 'R1', 'A01', 3, '方管'),
    ('1.4',     '14_原材料', '14ST10001', 'A01', 8, '自攻钉'),
    ('1.5',     '14_原材料', '14CM01007.04', 'A01', 1, '子零件'),
    ('1.5.1',   '14_原材料', 'R8', 'A01', 1, '子零件下级'),
]


@pytest.fixture
def bom_app():
    app = Flask(__name__)
    app.config.update(SQLALCHEMY_DATABASE_URI='sqlite://', SQLALCHEMY_TRACK_MODIFICATIONS=False)
    db.init_app(app)
    tables = [ImportProductRaw.__table__, MaterialBom.__table__, MaterialBomLine.__table__,
              ErpCodeRule.__table__, ErpGroupCategory.__table__, ProductMaterial.__table__,
              MaterialDisableKeyword.__table__, MaterialImage.__table__, ProductFinished.__table__,
              CostSnapshot.__table__, CostSnapshotSku.__table__, CostBomNode.__table__,
              CostMaterialPrice.__table__]

    def reset_caches():
        material_service.invalidate_rule_cache()
        material_service.invalidate_group_config_cache()
        material_service.invalidate_override_cache()

    with app.app_context():
        db.metadata.create_all(bind=db.engine, tables=tables)
        reset_caches()
        # ERP：成品/产成品只到 -A，半成品/原材料带 -A01；S1 在 ERP 里没有
        for code, name in [('F1-A', 'ERP书桌'), ('P1-A', 'ERP桌面'),
                           ('M1-A01', 'ERP钢架'), ('R1-A01', 'ERP方管')]:
            db.session.add(ImportProductRaw(code=code, name=name, group_code='G',
                                            group_name='组', imported_at=now_cst()))
        db.session.commit()
        yield app
        db.session.remove()
        reset_caches()
        db.metadata.drop_all(bind=db.engine, tables=list(reversed(tables)))


# ── 解析与拆分 ────────────────────────────────────────

def test_parse_skips_14st10_and_dot_subparts_with_subtree():
    rows, skipped = parse_bom_rows(_pdm(SAMPLE))
    assert skipped == 3                          # 14ST10001、14CM01007.04 及其下级 R8
    assert [r['code'] for r in rows] == ['F1', 'P1', 'M1', 'R1', 'S1', 'M1', 'R1']
    assert rows[0]['category'] == '成品' and rows[0]['row'] == 2
    boms = build_single_level_boms(rows)
    assert set(boms) == {('F1', 'A01'), ('P1', 'A01'), ('M1', 'A01')}
    assert [(l['code'], l['qty']) for l in boms[('F1', 'A01')]['lines']] == [
        ('P1', 1), ('S1', 4), ('M1', 1)]
    assert [(l['code'], l['qty']) for l in boms[('M1', 'A01')]['lines']] == [('R1', 3)]


def test_codes_and_versions_are_uppercased():
    rows, _ = parse_bom_rows(_pdm([('1', '11_成品', 'f1', 'a01', 1, '书桌'),
                                   ('1.1', '14_原材料', 'r1', 'a01', 2, '方管')]))
    assert [(r['code'], r['version']) for r in rows] == [('F1', 'A01'), ('R1', 'A01')]


def test_erp_format_parses_drawing_with_and_without_version():
    boms = _split(_erp([
        ('1', 'F1-A01', '书桌', '1.2米', 1, 'PCS'),
        ('1.1', 'R1-A01', '方管', 'Φ20', 2.5, 'PCS'),
        ('1.2', '170401001', '赠品', None, 1, 'PCS'),
    ]))
    lines = boms[('F1', 'A01')]['lines']
    assert [(l['code'], l['version'], l['qty']) for l in lines] == [
        ('R1', 'A01', 2.5), ('170401001', '', 1)]
    assert boms[('F1', 'A01')]['head']['spec'] == '1.2米'


@pytest.mark.parametrize('level, message', [
    (1.1, '数字格式'),
    ('1..1', '格式不正确'),
])
def test_bad_level_cell_rejected_with_row_number(level, message):
    with pytest.raises(UploadValidationError, match=f'第 3 行.*{message}'):
        parse_bom_rows(_pdm([('1', '11_成品', 'F1', 'A01', 1, '书桌'),
                             (level, '14_原材料', 'R1', 'A01', 1, '方管')]))


def test_integer_root_level_and_text_1_10_are_accepted():
    rows = [('1', '11_成品', 'F1', 'A01', 1, '书桌')] + [
        (f'1.{i}', '14_原材料', f'R{i}', 'A01', 1, 'x') for i in range(1, 11)]
    rows[0] = (1, *rows[0][1:])                  # 根层次写成整数单元格
    boms = _split(_pdm(rows))
    assert len(boms[('F1', 'A01')]['lines']) == 10   # 1.1 与 1.10 是两个不同子件


@pytest.mark.parametrize('rows, message', [
    ([('1', '11_成品', 'F1', 'A01', 1, 'a'), ('1.1', '14_原材料', 'R1', 'A01', 1, 'b'),
      ('1.1', '14_原材料', 'R2', 'A01', 1, 'c')], '第 4 行与第 3 行层次重复'),
    ([('1', '11_成品', 'F1', 'A01', 1, 'a'), ('1.2.1', '14_原材料', 'R1', 'A01', 1, 'b')],
     '第 3 行.*找不到上级 1.2'),
    ([('1.1', '14_原材料', 'R1', 'A01', 1, 'b'), ('1', '11_成品', 'F1', 'A01', 1, 'a')],
     '第 2 行.*找不到上级 1'),
])
def test_structure_errors_rejected(rows, message):
    with pytest.raises(UploadValidationError, match=message):
        _split(_pdm(rows))


def test_inconsistent_repeated_expansion_rejected():
    rows = SAMPLE[:6] + [('1.3.1', '14_原材料', 'R9', 'A01', 9, '不一致')]
    with pytest.raises(UploadValidationError, match='M1-A01 在第 4 行和第 7 行各展开了一次'):
        _split(_pdm(rows))


@pytest.mark.parametrize('qty, message', [
    (None, '「」无效'), ('abc', '「abc」无效'), (0, '必须大于 0'),
    (-1, '必须大于 0'), (1.23456, '小数超过 4 位'),
])
def test_invalid_quantity_rejected(qty, message):
    with pytest.raises(UploadValidationError, match=f'第 3 行数量.*{message}'):
        parse_bom_rows(_pdm([('1', '11_成品', 'F1', 'A01', 1, '书桌'),
                             ('1.1', '14_原材料', 'R1', 'A01', qty, '方管')]))


def test_duplicate_child_lines_sum_only_with_same_unit():
    base = [('1', '11_成品', 'F1', 'A01', 1, '书桌')]
    same = _split(_pdm(base + [('1.1', '14_原材料', 'R1', 'A01', 2, '方管'),
                               ('1.2', '14_原材料', 'R1', 'A01', 3, '方管')]))
    assert same[('F1', 'A01')]['lines'][0]['qty'] == 5
    with pytest.raises(UploadValidationError, match='单位不同'):
        _split(_pdm(base + [('1.1', '14_原材料', 'R1', 'A01', 2, '方管', 'PCS'),
                            ('1.2', '14_原材料', 'R1', 'A01', 3, '方管', 'SET')]))


def test_field_too_long_rejected():
    with pytest.raises(UploadValidationError, match='第 3 行编码超过 64 个字符'):
        parse_bom_rows(_pdm([('1', '11_成品', 'F1', 'A01', 1, '书桌'),
                             ('1.1', '14_原材料', 'R' * 65, 'A01', 1, '方管')]))


# ── 导入 / 覆盖 / 版本 ─────────────────────────────────

def test_import_maps_erp_codes_and_overwrites_same_version(bom_app):
    with bom_app.app_context():
        res = material_bom_service.import_file(_pdm(SAMPLE), 'a.xlsx', 'tester')
        assert res.success, res.message
        assert res.data['created'] == 3 and res.data['updated'] == 0
        assert res.data['skipped'] == 3
        assert res.data['roots'] == ['F1-A01']
        assert res.data['unmatched'] == ['S1-A01']
        f1 = MaterialBom.query.filter_by(code='F1', version='A01').one()
        assert f1.erp_code == 'F1-A'            # A01 → 退到只有字母的 -A
        assert MaterialBom.query.filter_by(code='M1').one().erp_code == 'M1-A01'

        # 同一研发版本再导一次（小写写法也算同一版本）：覆盖子件，不新增表头
        changed = [('1', '11_成品', 'f1', 'a01', 1, '书桌')] + [
            row if row[2] != 'S1' else ('1.2', '14_原材料', 'S1', 'A01', 6, '螺钉')
            for row in SAMPLE[1:]]
        res2 = material_bom_service.import_file(_pdm(changed), 'b.xlsx', 'tester')
        assert res2.data['created'] == 0 and res2.data['updated'] == 3
        assert MaterialBom.query.count() == 3
        assert MaterialBomLine.query.filter_by(bom_id=f1.id, code='S1').one().qty == 6

        # 新研发版本 A02 → 同一个 ERP 物料 F1-A 下多一个版本
        a02 = [(l, c, code, 'A02' if code == 'F1' else v, q, d)
               for l, c, code, v, q, d in SAMPLE]
        material_bom_service.import_file(_pdm(a02), 'c.xlsx', 'tester')
        # 卡片列出该物料下的全部研发 BOM（新→旧），各带下级数量
        view = material_bom_service.for_material('F1-A').data
        assert [(v['drawing'], v['line_count']) for v in view['versions']] == [
            ('F1-A02', 3), ('F1-A01', 3)]
        assert 'tree' not in view


def test_invalid_file_writes_nothing(bom_app):
    with bom_app.app_context():
        bad = SAMPLE[:4] + [('1.2', '14_原材料', 'S1', 'A01', 'x', '螺钉')]
        res = material_bom_service.import_file(_pdm(bad), 'a.xlsx', 'tester')
        assert not res.success and '第 6 行数量' in res.message
        assert MaterialBom.query.count() == 0 and MaterialBomLine.query.count() == 0


def test_import_rejects_unknown_format(bom_app):
    wb = openpyxl.Workbook()
    wb.active.append(['编号', '名字'])
    wb.active.append(['A', 'B'])
    out = io.BytesIO()
    wb.save(out)
    with bom_app.app_context():
        res = material_bom_service.import_file(out.getvalue(), 'x.xlsx', 'tester')
        assert not res.success and '无法识别' in res.message


# ── 展开 / 反查 ───────────────────────────────────────

def test_tree_expands_all_levels_with_erp_names(bom_app):
    with bom_app.app_context():
        material_bom_service.import_file(_pdm(SAMPLE), 'a.xlsx', 'tester')
        f1 = MaterialBom.query.filter_by(code='F1').one()
        data = material_bom_service.tree(f1.id).data
        assert data['bom']['name'] == 'ERP书桌'
        top = data['children']
        assert [n['drawing'] for n in top] == ['P1-A01', 'S1-A01', 'M1-A01']
        p1 = top[0]
        assert p1['name'] == 'ERP桌面' and p1['erp_code'] == 'P1-A'
        assert p1['children'][0]['drawing'] == 'M1-A01'
        assert p1['children'][0]['children'][0]['name'] == 'ERP方管'
        assert 'children' not in top[1]          # 螺钉没有下级
        assert top[1]['erp_code'] is None and top[1]['name'] == '螺钉'


def test_where_used_direct_parents_and_top_products(bom_app):
    with bom_app.app_context():
        material_bom_service.import_file(_pdm(SAMPLE), 'a.xlsx', 'tester')
        r1 = material_bom_service.for_material('R1-A01').data
        assert r1['versions'] == []
        assert [(d['drawing'], d['qty']) for d in r1['direct_parents']] == [('M1-A01', 3)]
        assert [t['drawing'] for t in r1['top_products']] == ['F1-A01']
        m1 = material_bom_service.for_material('M1-A01').data
        assert sorted(d['drawing'] for d in m1['direct_parents']) == ['F1-A01', 'P1-A01']
        assert [t['drawing'] for t in m1['top_products']] == ['F1-A01']
        assert [(v['drawing'], v['line_count']) for v in m1['versions']] == [('M1-A01', 1)]
        f1 = material_bom_service.for_material('F1-A').data
        assert f1['direct_parents'] == [] and f1['top_products'] == []
        assert material_bom_service.codes_with_bom(['F1-A', 'R1-A01', 'M1-A01']) == {'F1-A', 'M1-A01'}


def test_where_used_merges_research_versions_of_same_erp_code(bom_app):
    """同一 ERP 物料 Q1-A 的两个研发版本分别被两个顶层引用：两个顶层都返回且不重复。"""
    with bom_app.app_context():
        db.session.add(ImportProductRaw(code='Q1-A', name='ERP件', group_code='G',
                                        group_name='组', imported_at=now_cst()))
        db.session.commit()
        material_bom_service.import_file(_pdm([
            ('1', '11_成品', 'T1', 'A01', 1, '甲'), ('1.1', '12_产成品', 'Q1', 'A01', 1, '件'),
            ('1.1.1', '14_原材料', 'R1', 'A01', 1, '方管')]), 'a.xlsx', 'tester')
        material_bom_service.import_file(_pdm([
            ('1', '11_成品', 'T2', 'A01', 1, '乙'), ('1.1', '12_产成品', 'Q1', 'A02', 2, '件'),
            ('1.1.1', '14_原材料', 'R1', 'A01', 1, '方管')]), 'b.xlsx', 'tester')
        q = material_bom_service.for_material('Q1-A').data
        assert [v['version'] for v in q['versions']] == ['A02', 'A01']
        assert sorted(d['drawing'] for d in q['direct_parents']) == ['T1-A01', 'T2-A01']
        assert [t['drawing'] for t in q['top_products']] == ['T1-A01', 'T2-A01']


# ── ERP 重导补关联 ────────────────────────────────────

def test_relink_fills_only_unmatched_after_erp_import(bom_app):
    with bom_app.app_context():
        material_bom_service.import_file(_pdm(SAMPLE), 'a.xlsx', 'tester')
        f1 = MaterialBom.query.filter_by(code='F1').one()
        s1 = MaterialBomLine.query.filter_by(bom_id=f1.id, code='S1').one()
        assert s1.erp_code is None
        # 已匹配的人为改成别的值，验证补关联不会重写它
        MaterialBomLine.query.filter_by(bom_id=f1.id, code='P1').update({'erp_code': 'KEEP'})
        db.session.add(ImportProductRaw(code='S1-A01', name='ERP螺钉', group_code='G',
                                        group_name='组', imported_at=now_cst()))
        db.session.commit()
        stats = material_bom_service.relink_unmatched_erp_codes()
        assert stats == {'bom_headers_relinked': 0, 'bom_lines_relinked': 1}
        assert db.session.get(MaterialBomLine, s1.id).erp_code == 'S1-A01'
        assert MaterialBomLine.query.filter_by(bom_id=f1.id, code='P1').one().erp_code == 'KEEP'
        assert material_bom_service.relink_unmatched_erp_codes() == {
            'bom_headers_relinked': 0, 'bom_lines_relinked': 0}


# ── 删除保护 ──────────────────────────────────────────

def test_delete_referenced_bom_requires_force(bom_app):
    with bom_app.app_context():
        material_bom_service.import_file(_pdm(SAMPLE), 'a.xlsx', 'tester')
        m1 = MaterialBom.query.filter_by(code='M1').one()
        res = material_bom_service.delete_bom(m1.id)
        assert not res.success and res.data['needs_force'] is True
        assert [r['drawing'] for r in res.data['references']] == ['F1-A01', 'P1-A01']
        assert db.session.get(MaterialBom, m1.id) is not None

        assert material_bom_service.delete_bom(m1.id, force=True).success
        assert db.session.get(MaterialBom, m1.id) is None
        # 上级里 M1 仍在，只是不再可展开
        f1 = MaterialBom.query.filter_by(code='F1').one()
        m1_node = [n for n in material_bom_service.tree(f1.id).data['children']
                   if n['code'] == 'M1'][0]
        assert m1_node['bom_id'] is None and 'children' not in m1_node

        # 顶层成品没有被引用，直接可删
        assert material_bom_service.delete_bom(f1.id).success


def test_bom_routes_do_not_collide_with_item_routes():
    app = Flask(__name__)
    app.register_blueprint(material_bp, url_prefix='/api/material')
    routes = app.url_map.bind('localhost')
    assert routes.match('/api/material/items/A/B/bom') == (
        'material.material_item_bom', {'code': 'A/B'})
    assert routes.match('/api/material/items/A/B')[0] == 'material.material_detail'
    assert routes.match('/api/material/boms/3/tree')[0] == 'material.material_bom_tree'
    assert routes.match('/api/material/boms/import', method='POST')[0] == 'material.import_material_bom'


def test_list_groups_boms_by_material_type(bom_app):
    """左侧列表按物料类型分类：与物料表同一套判定；对应不到 ERP 的归「未匹配」。"""
    with bom_app.app_context():
        for group, flag in [('GF', 'is_finished'), ('GP', 'is_packaged'), ('GS', 'is_semi')]:
            db.session.add(ErpGroupCategory(group_code=group, **{flag: True}))
        for code, group in [('F1-A', 'GF'), ('P1-A', 'GP'), ('M1-A01', 'GS')]:
            ImportProductRaw.query.filter_by(code=code).update({'group_code': group})
        db.session.commit()
        material_service.invalidate_group_config_cache()
        material_bom_service.import_file(_pdm(SAMPLE), 'a.xlsx', 'tester')
        # ERP 里没有的父件：归入 unmatched
        material_bom_service.import_file(_pdm([
            ('1', '11_成品', 'Z9', 'A01', 1, '无'), ('1.1', '14_原材料', 'R1', 'A01', 1, '方管')]),
            'b.xlsx', 'tester')

        data = material_bom_service.list_boms().data
        assert data['all_total'] == data['total'] == 4
        assert data['type_counts']['finished'] == 1
        assert data['type_counts']['packaged'] == 1
        assert data['type_counts']['semi'] == 1
        assert data['type_counts']['unmatched'] == 1
        assert {i['drawing']: i['material_types'] for i in data['items']}['F1-A01'] == ['finished']

        semi = material_bom_service.list_boms(material_type='semi').data
        assert [i['drawing'] for i in semi['items']] == ['M1-A01']
        assert semi['total'] == 1 and semi['all_total'] == 4
        # 关键词与分类叠加；分类计数跟随关键词
        kw = material_bom_service.list_boms(keyword='Z9').data
        assert kw['all_total'] == 1 and kw['type_counts']['unmatched'] == 1
        assert kw['type_counts']['finished'] == 0


def test_export_xlsx_matches_tree_with_hierarchical_seq(bom_app):
    with bom_app.app_context():
        material_bom_service.import_file(_pdm(SAMPLE), 'a.xlsx', 'tester')
        f1 = MaterialBom.query.filter_by(code='F1').one()
        data, name = material_bom_service.export_xlsx(f1.id)
        assert name == 'BOM-F1-A01.xlsx'
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        assert ws['A1'].value.startswith('BOM：F1-A01')
        assert [c.value for c in ws[4]] == ['序号', '层级', '图纸编码', 'ERP编码', '名称', '数量', '单位']
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=5)]
        # 与页面树一致：P1 → M1 → R1，S1，M1 → R1
        assert [(r[0], r[1], r[2]) for r in rows] == [
            ('1', 1, 'P1-A01'), ('1.1', 2, 'M1-A01'), ('1.1.1', 3, 'R1-A01'),
            ('2', 1, 'S1-A01'), ('3', 1, 'M1-A01'), ('3.1', 2, 'R1-A01'),
        ]
        assert rows[0][3] == 'P1-A' and rows[0][4] == 'ERP桌面'
        assert rows[3][3] is None and rows[3][4] == '螺钉'   # ERP 未匹配：ERP编码留空，名称用文件里的
        assert rows[2][5] == 3 and rows[2][6] == 'PCS'
        assert material_bom_service.export_xlsx(99999)[0] is None


def test_export_route_is_registered():
    app = Flask(__name__)
    app.register_blueprint(material_bp, url_prefix='/api/material')
    routes = app.url_map.bind('localhost')
    assert routes.match('/api/material/boms/3/export')[0] == 'material.export_material_bom'


# ── 实时计价 ──────────────────────────────────────────

def _batch(order_no, day, finished_codes=None):
    """模拟一次采购导入批次；finished_codes=None 表示老批次（没有记录订单成品）。"""
    batch = CostSnapshot(order_no=order_no, snapshot_date=day, notes='采购导入价格')
    db.session.add(batch)
    db.session.flush()
    for code in finished_codes or []:
        db.session.add(CostSnapshotSku(snapshot_id=batch.id, finished_code=code))
    db.session.commit()
    return batch.id


def _price(code, price, day, source='bom_import', batch_id=None):
    node = CostBomNode.query.filter_by(code=code).first()
    if node is None:
        node = CostBomNode(code=code, code_with_version=f'{code}-A01', node_type='material')
        db.session.add(node)
        db.session.flush()
    db.session.add(CostMaterialPrice(node_id=node.id, unit_price=price, price_date=day, source=source,
                                     snapshot_id=batch_id))
    db.session.commit()


def _flat(nodes):
    for n in nodes:
        yield n
        yield from _flat(n.get('children') or [])


def _add_erp_raw(code):
    db.session.add(ImportProductRaw(code=code, name=f'ERP{code}', group_code='G', group_name='组',
                                    imported_at=now_cst()))
    db.session.commit()


def test_tree_prices_only_when_children_complete(bom_app):
    """齐全才计价：下级有一项没价格，部件就是未计价（不给残缺合计），只给缺价数。"""
    from datetime import date
    with bom_app.app_context():
        material_bom_service.import_file(_pdm(SAMPLE), 'a.xlsx', 'tester')
        f1 = MaterialBom.query.filter_by(code='F1').one()
        # 没有价格权限时不带任何价格字段
        plain = material_bom_service.tree(f1.id).data
        assert 'unit_price' not in plain['bom'] and 'unit_price' not in plain['children'][0]

        none = material_bom_service.tree(f1.id, include_price=True).data
        # 缺价按原材料种类计：R1 出现两处只算一种 → R1、S1 共 2 种
        assert none['bom']['unit_price'] is None and none['bom']['missing'] == 2

        _price('R1', 3, date(2024, 1, 1))
        _price('R1', 4, date(2024, 6, 1))
        _price('P1', 999, date(2024, 1, 1), source='bom_calc')   # 旧推算价不参与
        partial = material_bom_service.tree(f1.id, include_price=True).data
        p1, s1, m1 = partial['children']
        assert p1['unit_price'] == 24 and p1['price_source'] == 'calc'      # P1 下级齐全：2 × (3 × 4)
        assert m1['unit_price'] == 12
        assert s1['unit_price'] is None and s1['missing'] == 1
        # 根：S1 没价格 → 未计价，不给 36 这种残缺合计
        assert partial['bom']['unit_price'] is None and partial['bom']['missing'] == 1

        _price('S1', 0.5, date(2024, 6, 1))
        full = material_bom_service.tree(f1.id, include_price=True).data
        assert full['bom']['unit_price'] == 38.0 and full['bom']['missing'] == 0   # 24 + 4×0.5 + 12
        # 计价日期在 S1 有价之前：又回到未计价
        early = material_bom_service.tree(f1.id, include_price=True, price_date=date(2024, 3, 1)).data
        assert early['bom']['unit_price'] is None and early['children'][0]['unit_price'] == 18


def test_own_price_and_no_price_flag_make_items_complete(bom_app):
    from datetime import date
    from services.product.material import material_service
    with bom_app.app_context():
        _add_erp_raw('S1-A01')                          # 让 S1 能对应到 ERP 物料，才能打「不计价」标记
        material_bom_service.import_file(_pdm(SAMPLE), 'a.xlsx', 'tester')
        f1 = MaterialBom.query.filter_by(code='F1').one()
        _price('M1', 50, date(2024, 1, 1))             # 外购半成品：R1 无价格时用自身外购价
        data = material_bom_service.tree(f1.id, include_price=True).data
        m1 = data['children'][2]
        assert m1['unit_price'] == 50 and m1['price_source'] == 'own'
        assert data['children'][0]['unit_price'] == 100
        assert data['bom']['unit_price'] is None and data['bom']['missing'] == 1   # S1 仍缺

        assert not material_service.save_item('S1-A01', {'no_price': 'yes'}).success
        assert material_service.save_item('S1-A01', {'no_price': True}).data['no_price'] is True
        data = material_bom_service.tree(f1.id, include_price=True).data
        assert data['children'][1]['price_source'] == 'free' and data['children'][1]['unit_price'] == 0
        assert data['bom']['unit_price'] == 150.0 and data['bom']['missing'] == 0


def test_cost_overview_starts_when_complete(bom_app):
    from datetime import date
    with bom_app.app_context():
        material_bom_service.import_file(_pdm(SAMPLE), 'a.xlsx', 'tester')
        # 订单 A：本产品 F1 的订单，但 S1 没价格 → 还没开始计价
        a = _batch('A-20240101', date(2024, 1, 1), ['F1-A'])
        _price('R1', 3, date(2024, 1, 1), batch_id=a)
        data = material_bom_service.calc_price('F1-A').data
        assert data['current']['unit_price'] is None
        assert (data['current']['priced'], data['current']['total'], data['current']['missing']) == (1, 2, 1)
        assert data['started'] is None and data['history'] == [] and data['composition'] == []
        assert [(m['drawing'], m['qty'], m['parents']) for m in data['missing_items']] == [('S1-A01', 4.0, ['F1-A01'])]

        # 订单 B（别的产品 Y1）带来了 S1 的价格 → 这一刻齐全，开始计价
        b = _batch('B-20240601', date(2024, 6, 1), ['Y1-A01'])
        _price('R1', 4, date(2024, 6, 1), batch_id=b)
        _price('S1', 0.5, date(2024, 6, 1), batch_id=b)
        # 订单 C：本产品的订单，R1 涨到 5
        c = _batch('C-20240701', date(2024, 7, 1), ['F1-A01'])
        _price('R1', 5, date(2024, 7, 1), batch_id=c)

        data = material_bom_service.calc_price('F1-A').data
        assert data['started'] == {'date': '2024-06-01', 'order_no': 'B-20240601'}
        assert data['current']['unit_price'] == 47.0 and data['current']['missing'] == 0   # 9×5 + 4×0.5
        assert data['missing_items'] == []
        assert [(x['drawing'], x['amount']) for x in data['composition']] == [
            ('P1-A01', 30.0), ('M1-A01', 15.0), ('S1-A01', 2.0)]
        # 历史只从开始计价那一单往后；订单 A 不出现
        assert [(h['order_no'], h['unit_price'], h['related'], h['delta']) for h in data['history']] == [
            ('C-20240701', 47.0, True, 9.0), ('B-20240601', 38.0, False, None)]

        # 半成品 M1 在订单 A 那天就齐全（下级只有 R1），开始得更早
        m1 = material_bom_service.calc_price('M1-A01').data
        assert m1['started'] == {'date': '2024-01-01', 'order_no': 'A-20240101'}
        assert [h['related'] for h in m1['history']] == [True, False, True]


def test_export_includes_price_columns_only_with_permission(bom_app):
    from datetime import date
    with bom_app.app_context():
        material_bom_service.import_file(_pdm(SAMPLE), 'a.xlsx', 'tester')
        _price('R1', 4, date(2024, 6, 1))
        _price('S1', 0.5, date(2024, 6, 1))
        f1 = MaterialBom.query.filter_by(code='F1').one()
        data, _ = material_bom_service.export_xlsx(f1.id, include_price=True)
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        assert [c.value for c in ws[4]][-3:] == ['单价', '金额', '价格日期']
        assert '合计单价：38' in ws['A2'].value
        first = [c.value for c in ws[5]]
        assert first[7] == 24 and first[8] == 24
        plain, _ = material_bom_service.export_xlsx(f1.id)
        assert [c.value for c in openpyxl.load_workbook(io.BytesIO(plain)).active[4]][-1] == '单位'


def test_calc_price_route_registered():
    from routes.product.material import material_cost_bp
    app = Flask(__name__)
    app.register_blueprint(material_bp, url_prefix='/api/material')
    app.register_blueprint(material_cost_bp, url_prefix='/api/material')
    routes = app.url_map.bind('localhost')
    assert routes.match('/api/material/items/A/B/calc-price') == (
        'material_cost.material_calc_price', {'code': 'A/B'})


def test_tree_and_card_missing_counts_agree(bom_app):
    """BOM 树合计上的「缺N」和卡片「已计价 N / M 种」按同一口径（原材料种类）。"""
    from datetime import date
    with bom_app.app_context():
        material_bom_service.import_file(_pdm(SAMPLE), 'a.xlsx', 'tester')
        _price('S1', 0.5, date(2024, 6, 1))          # 只剩 R1（出现在两处）缺价
        f1 = MaterialBom.query.filter_by(code='F1').one()
        tree = material_bom_service.tree(f1.id, include_price=True).data
        card = material_bom_service.calc_price('F1-A').data['current']
        assert tree['bom']['missing'] == card['missing'] == 1
        assert (card['priced'], card['total']) == (1, 2)


def test_assemblies_cannot_be_no_price(bom_app):
    from services.product.material import material_service
    with bom_app.app_context():
        db.session.add(ErpGroupCategory(group_code='GS', is_semi=True))
        db.session.add(ImportProductRaw(code='SEMI-A01', name='半成品', group_code='GS', group_name='半成品',
                                        imported_at=now_cst()))
        _add_erp_raw('RAW-A01')
        material_service.invalidate_group_config_cache()
        res = material_service.save_item('SEMI-A01', {'no_price': True})
        assert not res.success and '不能设为不计价' in res.message
        # 原材料可以；之后被单独指定成半成品时，不计价标记自动取消
        assert material_service.save_item('RAW-A01', {'no_price': True}).data['no_price'] is True
        changed = material_service.save_item('RAW-A01', {'type_override': ['semi']}).data
        assert changed['no_price'] is False
        material_service.invalidate_override_cache()

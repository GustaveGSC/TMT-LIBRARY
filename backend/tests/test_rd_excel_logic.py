"""RD Excel 纯逻辑的行为护栏。"""

import io

import pytest
from openpyxl import Workbook, load_workbook

import routes.rd as rd_routes
import services.rd.change_documents as change_documents
from services.rd.change_documents import (
    _derive_new_drawing,
    _parse_bom,
    build_ecr_xlsx,
    build_ecn_xlsx,
    compare_bom,
    parse_ecr_rows_xlsx,
)
from services.rd.pdm_to_bom import build_bom_data, build_erp_data
from upload_validation import UploadValidationError


def _write_bom(path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['层次', '图号', '品名', '规格', '数量', '单位', '状态'])
    for row in rows:
        sheet.append(row)
    workbook.save(path)


_PDM_HEADERS = [
    '层次', '物料编码', '版本', '一级分类', '二级分类', '三级分类', '描述',
    '数量', '单位', '状态', '适配产品', '规格', '表面处理', '备注', '颜色',
    '系列版本', '类别', '尺寸', '材料',
]


def _write_pdm_bom(path, rows, headers=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers or _PDM_HEADERS)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_validate_bom_reports_open_failure(monkeypatch, tmp_path):
    monkeypatch.setattr(
        change_documents, 'report_internal_error',
        lambda context: 'fixed-error-id' if context == '变更前文件读取失败' else None,
    )

    message = change_documents.validate_bom(tmp_path / 'missing.xlsx', role='before')

    assert message == '变更前文件无法读取（错误编号：fixed-error-id）'


def test_compare_bom_characterizes_version_change(tmp_path):
    assert not hasattr(rd_routes, '_compare_bom')
    assert not hasattr(rd_routes, '_validate_bom')

    before_path = tmp_path / 'before.xlsx'
    after_path = tmp_path / 'after.xlsx'
    root = ['1', 'ROOT-A01', '整机', 'ROOT_A01', 1, 'PCS', '已发布']
    _write_bom(before_path, [
        root,
        ['1.1', 'PART-A01', '零件', 'PART_A01', 2, 'PCS', '已发布'],
    ])
    _write_bom(after_path, [
        root,
        ['1.1', 'PART-A01', '零件', 'PART_A01', 2, 'PCS', '通用变更审核中'],
    ])

    result = compare_bom(before_path, after_path)

    assert result['stats'] == {'version': 1, 'added': 0, 'deleted': 0, 'total': 2}
    assert result['changes'] == [
        {
            'row_type': 'cancel', 'change_kind': '通用变更', 'level': '1.1',
            'main_drawing': 'ROOT-A01', 'drawing': 'PART-A01', 'name': '零件',
            'spec': 'PART_A01', 'qty': 2.0, 'change_method': '取消', 'seq': 1,
        },
        {
            'row_type': 'add', 'change_kind': '通用变更', 'level': '1.1',
            'main_drawing': 'ROOT-A01', 'drawing': 'PART-A02', 'name': '零件',
            'spec': 'PART_A02', 'qty': 2.0, 'change_method': '新增', 'seq': 2,
        },
    ]


def test_pdm_bom_parses_new_material_and_compare_assigns_a01(tmp_path):
    before_path = tmp_path / 'before.xlsx'
    after_path = tmp_path / 'after.xlsx'
    _write_bom(before_path, [
        ['1', 'ROOT-A01', '整机', 'ROOT_A01', 1, 'PCS', '已发布'],
    ])
    _write_pdm_bom(after_path, [
        ['1', 'ROOT', 'A01', '12_产成品', '01_桌类', 'XG_学习工场', '桌面',
         1, 'PCS', '已发布', '', '', '', '', '灰色', 'V2.3', '手摇', '1.2米', '榉木'],
        ['1.1', 'NEW001', '', '14_原材料', 'WD_木器', '04_榉木', '活动桌面板',
         2, 'PCS', '审核中', 'π桌', '1200x600', '喷粉', '', '', '', '', '', ''],
    ])

    parsed = _parse_bom(after_path)
    new_item = parsed[('ROOT', 'NEW001')]
    result = compare_bom(before_path, after_path)

    assert new_item['drawing'] == ''
    assert new_item['name'] == '原材料_木器_榉木_活动桌面板'
    assert new_item['spec'] == 'π桌_1200x600_喷粉_A01'
    added = next(item for item in result['changes'] if item['drawing'] == 'NEW001-A01')
    assert added['row_type'] == 'added'
    assert added['spec'] == 'π桌_1200x600_喷粉_A01'


def test_pdm_packaged_spec_and_three_status_versions(tmp_path):
    path = tmp_path / 'pdm.xlsx'
    _write_pdm_bom(path, [[
        '1', 'PACK', 'A02', '12_产成品', '01_桌类', 'XG_学习工场', '桌面',
        1, 'PCS', '通用变更审核中', '', '', '', '倾斜款', '木色', 'V2.3',
        '手摇', '1.2米', '榉木',
    ]])
    item = next(iter(_parse_bom(path).values()))
    assert item['name'] == '产成品_桌类_学习工场_桌面'
    assert item['spec'] == '(V2.3)倾斜款手摇1.2米榉木木色_A'
    assert _derive_new_drawing({'status': '审核中', 'code': 'N', 'drawing': '', 'version': ''}) == 'N-A01'
    assert _derive_new_drawing({'status': '通用变更审核中', 'code': 'P', 'drawing': 'P-A01', 'version': 'A01'}) == 'P-A02'
    assert _derive_new_drawing({'status': '非通用变更审核中', 'code': 'P', 'drawing': 'P-A01', 'version': 'A01'}) == 'P-B01'


def test_pdm_duplicate_required_header_and_unknown_status_are_rejected(tmp_path):
    duplicate_path = tmp_path / 'duplicate.xlsx'
    duplicate_headers = [*_PDM_HEADERS, '版本']
    _write_pdm_bom(duplicate_path, [], headers=duplicate_headers)
    error = change_documents.validate_bom(duplicate_path, role='any')
    assert '「版本」出现了 2 次' in error
    assert '第 3、20 列' in error

    unknown_path = tmp_path / 'unknown.xlsx'
    _write_pdm_bom(unknown_path, [[
        '1', 'CODE', 'A01', '14_原材料', 'WD_木器', '', '桌面',
        1, 'PCS', '草稿', '', '', '', '', '', '', '', '', '',
    ]])
    assert '未知状态：草稿' in change_documents.validate_bom(unknown_path, role='any')


def test_pdm_numeric_code_preserves_zero_number_format(tmp_path):
    path = tmp_path / 'numeric.xlsx'
    _write_pdm_bom(path, [[
        '1', 1234, 'A01', '14_原材料', 'WD_木器', '', '桌面',
        1, 'PCS', '已发布', '', '', '', '', '', '', '', '', '',
    ]])
    workbook = load_workbook(path)
    workbook.active.cell(2, 2).number_format = '000000'
    workbook.save(path)
    workbook.close()
    item = next(iter(_parse_bom(path).values()))
    assert item['code'] == '001234'


def test_bom_unknown_format_and_non_integer_numeric_code_are_rejected(tmp_path):
    unknown_path = tmp_path / 'unknown-format.xlsx'
    workbook = Workbook()
    workbook.active.append(['任意列', '状态'])
    workbook.active.append(['值', '已发布'])
    workbook.save(unknown_path)
    assert '无法识别的 BOM 文件格式' in change_documents.validate_bom(
        unknown_path, role='any',
    )

    numeric_path = tmp_path / 'non-integer-code.xlsx'
    _write_pdm_bom(numeric_path, [[
        '1', 1234.5, 'A01', '14_原材料', 'WD_木器', '', '桌面',
        1, 'PCS', '已发布', '', '', '', '', '', '', '', '', '',
    ]])
    with pytest.raises(UploadValidationError, match='非整数数值'):
        _parse_bom(numeric_path)


def test_ecr_xlsx_round_trip_preserves_key_fields_and_detail():
    assert not hasattr(rd_routes, '_build_ecr_xlsx')
    assert not hasattr(rd_routes, '_parse_ecr_rows_xlsx')
    assert not hasattr(rd_routes, '_parse_ecr_rows_xls')

    fields = {
        'issuing_unit': '研发',
        'date': '2026-07-23',
        'ecr_code': 'ECR-001',
        'project': '测试项目',
        'change_type': '设计变更',
        'distribution': ['研发', '采购'],
        'change_reason': '设计优化',
        'change_subject': '替换零件',
        'change_desc': '验证生成与解析行为',
        'submitter': 'tester',
    }
    changes = [{
        'seq': 1,
        'row_type': 'added',
        'level': '1.1',
        'main_drawing': 'ROOT-A01',
        'drawing': 'PART-A01',
        'name': '零件',
        'spec': '规格A',
        'change_method': '数量变更',
        'qty_desc': '1→2 PCS',
    }]

    workbook = load_workbook(io.BytesIO(build_ecr_xlsx(fields, changes)))
    parsed_fields, parsed_changes = parse_ecr_rows_xlsx(workbook.active)
    workbook.close()

    assert parsed_fields == fields
    assert parsed_changes == [{
        'seq': 1,
        'level': '1.1',
        'main_drawing': 'ROOT-A01',
        'drawing': 'PART-A01',
        'name': '零件',
        'spec': '规格A',
        'change_method': '数量变更',
        'change_kind': '',
        'substitution': '1→2 PCS',
        'handling': '',
        'responsible_person': '',
        'row_type': 'added',
    }]


def test_ecn_xlsx_preserves_document_and_detail_layout():
    assert not hasattr(rd_routes, '_build_ecn_xlsx')

    fields = {
        'issuing_unit': '研发',
        'product': 'TMT-01',
        'responsible': 'tester',
        'date': '2026-07-23',
        'ecn_code': 'ECN-001',
        'import_method': '立即导入',
        'distribution': ['研发'],
        'change_reason': '结构优化',
        'change_desc': 'ECN fixture',
        'affected_files': ['BOM'],
    }
    changes = [{
        'seq': 1,
        'row_type': 'added',
        'level': '1.1',
        'main_drawing': 'ROOT-A01',
        'drawing': 'PART-A02',
        'name': '零件',
        'spec': '规格B',
        'change_method': '新增',
        'change_kind': '通用变更',
        'handling': '立即导入',
        'responsible_person': 'tester',
    }]

    workbook = load_workbook(io.BytesIO(build_ecn_xlsx(fields, changes)))
    sheet = workbook.active

    assert sheet.title == '变更通知单'
    assert [sheet.cell(3, column).value for column in (2, 5, 9, 12, 13)] == [
        '研发', 'TMT-01', 'tester', '2026-07-23', 'ECN-001',
    ]
    assert [sheet.cell(10, column).value for column in (1, 2, 4, 6, 7, 9, 11, 12, 13, 14)] == [
        1, 'ROOT-A01', 'PART-A02', '1.1', '零件', '规格B',
        '新增', '通用变更', '立即导入', 'tester',
    ]
    assert '☑ 立即导入' in sheet.cell(4, 2).value
    assert '☑ BOM' in sheet.cell(8, 2).value
    workbook.close()


def test_pdm_builders_preserve_constants_and_parent_child_relations():
    assert not hasattr(rd_routes, '_ptb_build_erp_data')
    assert not hasattr(rd_routes, '_ptb_build_bom_data')

    columns = ['品号', '层次', '数量', '品名', '工作中心']
    table_data = [
        ['ROOT', '1', '1', '整机', '源工作中心'],
        ['CHILD', '1.1', '2', '子件', ''],
        ['GRAND', '1.1.1', '3', '孙件', ''],
        ['SIBLING', '1.2', '无效数量', '并列子件', ''],
    ]

    erp_rows = build_erp_data(columns, table_data)
    assert len(erp_rows) == 4
    assert len(erp_rows[0]) == 46
    assert erp_rows[0][1:3] == ['ROOT', '整机']
    assert erp_rows[0][16:18] == [1001, 100802]
    assert erp_rows[0][22] == 5101
    assert erp_rows[0][23] == 3101

    bom_rows = build_bom_data(columns, table_data, total_level=3)
    assert bom_rows == [
        [100, 'ROOT', 1, 'CHILD', 2, 1, '', '', '', '', '', 1, ''],
        [100, 'ROOT', 1, 'SIBLING', 1, 1, '', '', '', '', '', 1, ''],
        [100, 'CHILD', 1, 'GRAND', 3, 1, '', '', '', '', '', 1, ''],
    ]
